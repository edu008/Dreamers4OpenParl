from pathlib import Path

from openpyxl import Workbook

from parltopics.registry import load_registry
from parltopics.registry.aps_import import aps_slug, import_aps, refresh_descriptions


def _xlsx(path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "Domains"
    workbook["Domains"].append(["code", "domain_de", "domain_fr"])
    workbook["Domains"].append(["1", "Akteure", "Acteurs"])
    chapters = workbook.create_sheet("Chapters")
    chapters.append(["code", "chapter_de", "chapter_fr"])
    chapters.append(["1.1", "Grundfragen", "Questions"])
    topics = workbook.create_sheet("Topics")
    topics.append(["code", "topic_de", "topic_fr"])
    topics.append(["1.1.1", "Identität", "Identité"])
    workbook.save(path)


def test_import_aps_keeps_adhoc_and_is_idempotent(tmp_path: Path) -> None:
    xlsx = tmp_path / "aps.xlsx"
    _xlsx(xlsx)
    registry_file = tmp_path / "tags.yaml"
    registry_file.write_text(
        "registry_version: '2026-10-18.1'\n"
        "schemes:\n"
        "  aps: {label: APS, source: x, fixed: true}\n"
        "  adhoc: {label: Ad hoc, source: x, fixed: false}\n"
        "tags:\n"
        "  - tag_id: 'adhoc:a'\n    scheme: adhoc\n    status: active\n"
        "    pref_label: {de: A}\n    definition: {de: A}\n",
        encoding="utf-8",
    )
    count, version = import_aps(xlsx, registry_file)
    registry = load_registry(registry_file)
    assert count == 3
    assert [tag.tag_id for tag in registry.tags] == ["aps:1", "aps:1.1", "aps:1.1.1", "adhoc:a"]
    topic = registry.get("aps:1.1.1")
    assert topic.broader == ["aps:1.1"]
    assert topic.slug == "identitaet"
    assert topic.pref_label.fr == "Identité"
    assert registry.get("aps:1").broader == []
    assert registry.get("aps:1").definition.de.endswith("Umfasst: Grundfragen.")
    assert import_aps(xlsx, registry_file) == (3, version)


def test_descriptions_add_definition_and_synonyms_on_import_and_refresh(tmp_path: Path) -> None:
    xlsx = tmp_path / "aps.xlsx"
    _xlsx(xlsx)
    registry_file = tmp_path / "tags.yaml"
    registry_file.write_text(
        "registry_version: '2026-10-18.1'\n"
        "schemes:\n  aps: {label: APS, source: x, fixed: true}\n"
        "  adhoc: {label: Ad hoc, source: x, fixed: false}\ntags: []\n",
        encoding="utf-8",
    )
    _, plain = import_aps(xlsx, registry_file)
    (tmp_path / "aps_descriptions.yaml").write_text(
        "aps:1: {it: Attori}\n"
        "aps:1.1: {de: Grundsätze des Staates.}\n"
        "aps:1.1.1:\n  de: Selbstverständnis der Schweiz.\n"
        "  alt: {de: [Nationalfeiertag], fr: [fête nationale], it: []}\n",
        encoding="utf-8",
    )
    count, version = refresh_descriptions(registry_file)
    registry = load_registry(registry_file)
    assert (count, version != plain) == (3, True)
    assert registry.get("aps:1.1").definition.de == "Grundsätze des Staates. Umfasst: Identität."
    topic = registry.get("aps:1.1.1")
    assert topic.definition.de == "Selbstverständnis der Schweiz."
    assert topic.alt_labels.all() == ["Nationalfeiertag", "fête nationale"]
    assert registry.get("aps:1").pref_label.it == "Attori"
    assert registry.get("aps:1").definition.de.endswith("Umfasst: Grundfragen.")
    assert refresh_descriptions(registry_file) == (3, version)
    assert import_aps(xlsx, registry_file) == (3, version)


def test_aps_slug_drops_filler_words_and_umlauts() -> None:
    assert aps_slug("Verkehr und Kommunikation") == "verkehr-kommunikation"
    assert aps_slug("Infrastruktur und Umwelt") == "infrastruktur-umwelt"
    assert aps_slug("Europa- und Aussenpolitik") == "europa-aussenpolitik"
    assert aps_slug("Gesundheit, Sozialhilfe, Sport") == "gesundheit-sozialhilfe-sport"
    assert aps_slug("Identität") == "identitaet"
