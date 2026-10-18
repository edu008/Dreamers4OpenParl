"""APS-Schema aus aps_topics.xlsx in die Registry übernehmen (PRD F4, Anhang A)."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from parltopics.registry.loader import RegistryError

HEADER = (
    "# ParlTopics Tag-Registry: einzige Quelle aller Tags (PRD 7.2, Anhang A, ADR 0001).\n"
    "# APS-Einträge erzeugt `parltopics registry import-aps <aps_topics.xlsx>`.\n"
    "# Jede inhaltliche Änderung erhöht registry_version.\n"
    "# Neue Tags nur über Vorschläge und Review.\n"
    "# Schreibt ein Werkzeug die Datei neu, gehen diese Kommentare verloren (ADR 0001).\n"
)

SHEETS = {
    "Domains": ("code", "domain_de", "domain_fr"),
    "Chapters": ("code", "chapter_de", "chapter_fr"),
    "Topics": ("code", "topic_de", "topic_fr"),
}


def _rows(path: Path) -> dict[str, list[dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    result: dict[str, list[dict[str, str]]] = {}
    for sheet, columns in SHEETS.items():
        if sheet not in workbook.sheetnames:
            raise RegistryError(f"{path}: Blatt {sheet!r} fehlt")
        iterator = workbook[sheet].iter_rows(values_only=True)
        header = [str(cell) if cell is not None else "" for cell in next(iterator)]
        missing = [column for column in columns if column not in header]
        if missing:
            raise RegistryError(f"{path}: Blatt {sheet!r} ohne Spalten {', '.join(missing)}")
        rows = []
        for values in iterator:
            row = {
                key: ("" if value is None else str(value).strip())
                for key, value in zip(header, values, strict=False)
            }
            if row["code"]:
                rows.append(row)
        result[sheet] = rows
    return result


def _parent(code: str) -> str | None:
    return f"aps:{code.rsplit('.', 1)[0]}" if "." in code else None


def _description(path: list[tuple[str, str | None]], lang: str) -> str | None:
    """Beschreibung aus dem APS-Pfad. APS selbst liefert keine Beschreibungstexte."""
    names = [name for name, _ in path] if lang == "de" else [fr for _, fr in path]
    if any(name is None for name in names):
        return None
    if lang == "de":
        return f"Geschäfte zum APS-Thema «{names[-1]}» (Pfad: {' > '.join(names)})."
    return f"Objets relatifs au thème APS « {names[-1]} » (chemin : {' > '.join(names)})."


def _children_note(names: list[str], lang: str) -> str:
    """Kapitel und Bereiche haben in APS keinen Text: die Namen der Unterthemen ersetzen ihn."""
    joined = "; ".join(names)
    return f" Umfasst: {joined}." if lang == "de" else f" Comprend : {joined}."


def build_aps_tags(xlsx: Path) -> list[dict[str, Any]]:
    """Alle APS-Einträge (Bereich, Kapitel, Thema) als Registry-Dicts, sortiert nach Code."""
    sheets = _rows(xlsx)
    domains = {r["code"]: (r["domain_de"], r["domain_fr"]) for r in sheets["Domains"]}
    chapters = {r["code"]: (r["chapter_de"], r["chapter_fr"]) for r in sheets["Chapters"]}
    topics = {r["code"]: (r["topic_de"], r["topic_fr"]) for r in sheets["Topics"]}

    def entry(code: str, de: str, fr: str, path: list[tuple[str, str | None]]) -> dict[str, Any]:
        tag: dict[str, Any] = {
            "tag_id": f"aps:{code}",
            "scheme": "aps",
            "status": "active",
        }
        parent = _parent(code)
        if parent:
            tag["parent"] = parent
        tag["names"] = {"de": de, **({"fr": fr} if fr else {})}
        descriptions = {"de": _description(path, "de")}
        if all(fr_name for _, fr_name in path):
            descriptions["fr"] = _description(path, "fr")
        tag["descriptions"] = descriptions
        return tag

    def sort_key(code: str) -> list[int]:
        return [int(part) for part in code.split(".")]

    result = []
    for code in sorted([*domains, *chapters, *topics], key=sort_key):
        parts = code.split(".")
        path = [domains[parts[0]]]
        if len(parts) > 1:
            path.append(chapters[".".join(parts[:2])])
        if len(parts) > 2:
            path.append(topics[code])
        names = {1: domains, 2: chapters, 3: topics}[len(parts)][code]
        tag = entry(code, names[0], names[1], path)
        if len(parts) < 3:
            children = {1: chapters, 2: topics}[len(parts)]
            kids = [v for c, v in children.items() if c.startswith(f"{code}.")]
            for index, lang in enumerate(("de", "fr")):
                if lang in tag["descriptions"]:
                    tag["descriptions"][lang] += _children_note([k[index] for k in kids], lang)
        result.append(tag)
    return result


def import_aps(xlsx: Path, registry_path: Path) -> tuple[int, str]:
    """APS-Tags in die Registry schreiben, andere Schemas bleiben unverändert.

    Gibt (Anzahl APS-Tags, registry_version) zurück. Die Version steigt nur bei Änderung.
    """
    data = yaml.safe_load(registry_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise RegistryError(f"{registry_path}: erwartet ein YAML-Objekt auf oberster Ebene")
    aps_tags = build_aps_tags(xlsx)
    others = [tag for tag in data.get("tags", []) if tag.get("scheme") != "aps"]
    old_aps = [tag for tag in data.get("tags", []) if tag.get("scheme") == "aps"]
    if old_aps == aps_tags:
        return len(aps_tags), data["registry_version"]

    today = dt.date.today().isoformat()
    current_day, _, counter = data["registry_version"].partition(".")
    data["registry_version"] = f"{today}.{int(counter) + 1 if current_day == today else 1}"
    data["tags"] = [*aps_tags, *others]
    body = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, width=100)
    registry_path.write_text(HEADER + body, encoding="utf-8", newline="\n")
    return len(aps_tags), data["registry_version"]
